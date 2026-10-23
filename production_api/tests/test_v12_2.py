from resume_similarity import similarity,classify
from resume_journey import suggestions
from types import SimpleNamespace
from io import BytesIO

def test_near_duplicate_detected():
 a=SimpleNamespace(kind='project',title='Projeto de relatórios em Python',description='Criei relatórios automatizados para acompanhamento operacional usando Python e planilhas.')
 b=SimpleNamespace(kind='project',title='Projeto de relatorios em Python',description='Criei relatorios automatizados para acompanhamento operacional usando Python e planilhas')
 assert classify(a,b)=='duplicate'

def test_possible_duplicate_requires_review():
 a=SimpleNamespace(kind='project',title='Sistema de relatórios em Python',description='Criei relatorios automatizados para acompanhar vendas de clientes usando Python e planilhas.')
 b=SimpleNamespace(kind='project',title='Sistema de relatorios em Python',description='Criei relatorios automatizados para acompanhar vendas internas usando Python e planilhas.')
 assert classify(a,b)=='possible_duplicate'

def test_distinct_projects_remain_distinct():
 a=SimpleNamespace(kind='project',title='API de pagamentos',description='Implementei endpoint REST com testes de integração')
 b=SimpleNamespace(kind='project',title='Dashboard industrial',description='Monitorei produção com Excel e Power BI')
 assert classify(a,b) is None

def test_section_numbering():
 x=suggestions('1. EXPERIÊNCIA PROFISSIONAL\n- Analista 2023\nAtualizei indicadores.\n- Assistente 2024\nOrganizei planilhas.\n2. PROJETOS\n- Projeto Alpha\nCriei API.\n- Projeto Beta\nCriei relatórios.')
 assert len([e for e in x['evidence_drafts'] if e['kind']=='experience'])==2
 assert len([e for e in x['evidence_drafts'] if e['kind']=='project'])==2

def test_docx_table_preserves_order():
 from docx import Document
 from resume_ingestion import extract_resume
 d=Document();d.add_paragraph('EXPERIÊNCIA PROFISSIONAL');t=d.add_table(rows=1,cols=2);t.cell(0,0).text='Analista 2023';t.cell(0,1).text='Empresa A';d.add_paragraph('PROJETOS');d.add_paragraph('Projeto Beta - relatórios')
 f=BytesIO();d.save(f)
 x=extract_resume('resume.docx',f.getvalue())['text']
 assert x.index('EXPERIÊNCIA')<x.index('Analista 2023')<x.index('PROJETOS')
